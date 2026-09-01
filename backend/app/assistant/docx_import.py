"""Импорт .docx-шаблонов протокола (Блок 6): организация ведёт реальные бланки протоколов
допроса с DB mail-merge токенами вида `#{T1.PARTICIP_SURNAME}` — этот модуль их находит и
переписывает в Jinja-плейсхолдеры (`{{ T1_PARTICIP_SURNAME }}`), которые потом заполняет
`docgen.render()` через `docxtpl`.

Ключевой подвох Word: спелл-чекер/автозамена нередко режут один токен на несколько `<w:r>`-ранов
(например `#{T1.PARTIC` в одном ране и `IPANT}` в другом) — наивная замена по каждому рану в
отдельности такие токены просто не находит. Поэтому замена всегда идёт по СКЛЕЕННОМУ тексту
абзаца целиком (см. `_replace_in_paragraph`).
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Callable

from docx import Document

from ..models import Template

PLACEHOLDER_RE = re.compile(r"#\{([\w.]+)\}")
JINJA_VAR_RE = re.compile(r"\{\{\s*([\w]+)\s*\}\}")


def jinja_key(token: str) -> str:
    """Jinja-идентификаторы не допускают точку — `"T1.PARTICIP_SURNAME"` -> `"T1_PARTICIP_SURNAME"`."""
    return token.replace(".", "_")


def _iter_paragraphs(doc):
    """Все абзацы документа: тело, таблицы (рекурсивно — в ячейке может быть вложенная таблица),
    и колонтитулы каждой секции (простой проход по абзацам, без вложенных таблиц — в
    колонтитулах усложнять незачем)."""
    yield from doc.paragraphs
    yield from _iter_table_paragraphs(doc.tables)
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _iter_table_paragraphs(tables):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                yield from _iter_table_paragraphs(cell.tables)   # вложенные таблицы


def _paragraph_text(p) -> str:
    return "".join(run.text for run in p.runs)


def _replace_in_paragraph(p, transform: Callable[[str], str]) -> None:
    """Заменяет текст абзаца целиком через `transform(joined_text)`, а не по каждому рану —
    см. докстринг модуля про то, почему Word разбивает плейсхолдеры на несколько ранов."""
    runs = p.runs
    if not runs:
        return
    joined = "".join(run.text for run in runs)
    new_text = transform(joined)
    if new_text == joined:
        return
    runs[0].text = new_text          # форматирование первого рана сохраняется как есть
    for run in runs[1:]:
        run.text = ""


def scan_placeholders(path) -> list[str]:
    """Отсортированный список уникальных токенов `#{...}`, найденных во всём документе."""
    doc = Document(path)
    found: set[str] = set()
    for p in _iter_paragraphs(doc):
        found.update(PLACEHOLDER_RE.findall(_paragraph_text(p)))
    return sorted(found)


def scan_jinja_keys(path) -> set[str]:
    """Множество Jinja-переменных (`{{ VAR }}`, уже без точек), реально встречающихся в УЖЕ
    нормализованном докс-бланке. Нужно `docgen.render` как страховка: набор шагов шаблона в
    редакторе может разойтись с реальными плейсхолдерами файла (шаг удалили/переименовали
    плейсхолдер вручную с опечаткой — реальный случай, воспроизведённый пользователем) — тогда
    docxtpl тихо подставит пустую строку вместо `MISSING_MARKER`, и пропуск в документе останется
    незамеченным при вычитке. Сканирование самого файла, а не списка шагов, ловит это независимо
    от причины расхождения."""
    doc = Document(path)
    found: set[str] = set()
    for p in _iter_paragraphs(doc):
        found.update(JINJA_VAR_RE.findall(_paragraph_text(p)))
    return found


def uncovered_placeholders(tmpl: Template, docx_path) -> list[str]:
    """Плейсхолдеры, реально встречающиеся в `docx_path`, но не покрытые ни одним шагом `tmpl` и
    не выбранные как `qa_placeholder` (Блок 6) — см. `main.py::template_docx_coverage`. Отдельная
    функция без завязки на TemplateStore/config — проверяется без поднятия сервера/хранилища."""
    covered = {jinja_key(s.placeholder or s.key) for s in tmpl.steps}
    if tmpl.qa_placeholder:
        covered.add(jinja_key(tmpl.qa_placeholder))
    return sorted(scan_jinja_keys(docx_path) - covered)


def normalize_docx(src_path, dst_path) -> list[str]:
    """Переписывает все `#{TOKEN}` в `{{ TOKEN_с_подчёркиванием }}` и сохраняет результат в
    `dst_path`. `src_path` — что угодно, что принимает `docx.Document()` (путь или file-like,
    например `io.BytesIO` из загруженного файла). Возвращает отсортированный список уникальных
    ИСХОДНЫХ токенов (считаем их во время самого прохода замены, а не повторным сканированием —
    после замены токенов в тексте уже не остаётся)."""
    doc = Document(src_path)
    found: set[str] = set()

    def transform(text: str) -> str:
        def _sub(m: re.Match) -> str:
            token = m.group(1)
            found.add(token)
            return "{{ " + jinja_key(token) + " }}"
        return PLACEHOLDER_RE.sub(_sub, text)

    for p in _iter_paragraphs(doc):
        _replace_in_paragraph(p, transform)

    dst_path = Path(dst_path)
    dst_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst_path))
    return sorted(found)
